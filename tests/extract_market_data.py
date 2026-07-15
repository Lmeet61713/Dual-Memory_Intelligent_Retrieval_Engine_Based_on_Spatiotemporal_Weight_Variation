"""提取 xlsx 问卷数据和 docx 需求文档内容"""
import os
import json

# ====== 1. Read xlsx ======
print("=" * 80)
print("  XLSX 问卷数据")
print("=" * 80)

from openpyxl import load_workbook
xlsx_path = r"e:\mmmmmmmmmmmmmmmm\MyAgent\市场调研\359574280_按文本_智能物品记忆与定位系统需求调查问卷_97_97-1.xlsx"
wb = load_workbook(xlsx_path, data_only=True)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n--- Sheet: {sheet_name} ---")
    print(f"Rows: {ws.max_row}, Cols: {ws.max_column}")

    # Print headers (row 1)
    headers = []
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        headers.append(str(val) if val else "")
    print(f"Headers: {headers}")

    # Print all data rows
    for row_idx in range(2, ws.max_row + 1):
        vals = []
        for col_idx in range(1, ws.max_column + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            vals.append(str(val) if val is not None else "")
        print(f"  Row {row_idx}: {vals}")

wb.close()

# ====== 2. Read docx ======
print("\n" + "=" * 80)
print("  DOCX 需求文档")
print("=" * 80)

from docx import Document
docx_path = r"e:\mmmmmmmmmmmmmmmm\MyAgent\市场调研\智能物品记忆与定位系统需求调查.docx"
doc = Document(docx_path)

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        style = para.style.name if para.style else "Normal"
        print(f"[{style}] {text}")

# Also check tables
print("\n--- Tables ---")
for ti, table in enumerate(doc.tables):
    print(f"\nTable {ti+1}: {len(table.rows)} rows x {len(table.columns)} cols")
    for ri, row in enumerate(table.rows):
        cells = [cell.text.strip() for cell in row.cells]
        print(f"  Row {ri}: {cells}")
